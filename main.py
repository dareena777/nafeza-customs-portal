import io
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from flask import Flask, render_template
from fastapi import FastAPI, UploadFile, File, HTTPException, Response
from fastapi.middleware.cors import CORSMiddleware
from a2wsgi import WSGIMiddleware 

# Import custom OCR dependencies
from dpi import DPINormalizer
from ocr import extract_with_gemini, preprocess_file_bytes
from schema import InvoiceDataResponse


api_app = FastAPI(
    title="Elamry Logistics Intelligence API",
    description="Customs Document OCR & Reconciliation Engine",
    version="1.0.0"
)

api_app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

normalizer = DPINormalizer()

@api_app.post("/api/extract", response_model=InvoiceDataResponse)
async def extract_invoice(file: UploadFile = File(...)):
    """Async endpoint that handles image preprocessing and Gemini structured vision extraction."""
    try:
        contents = await file.read()
        mime_type = file.content_type or "application/pdf"
        
        # Preprocess document image / PDF bytes
        pages = preprocess_file_bytes(contents, mime_type, normalizer)
        
        api_key = os.environ.get("GEMINI_API_KEY")
        if not api_key:
            raise HTTPException(status_code=500, detail="GEMINI_API_KEY is not configured.")
            
        # Extract structured data using Gemini vision
        extracted_data = extract_with_gemini(pages, api_key)
        
        # Line-item reconciliation check
        if extracted_data.items:
            calc_sum = sum(
                (item.total_price or ((item.quantity or 0) * (item.unit_price or 0)))
                for item in extracted_data.items
            )
            extracted_data.calculated_total = round(calc_sum, 2)
            
            if extracted_data.bill_total is not None:
                extracted_data.totals_match = abs(
                    extracted_data.calculated_total - extracted_data.bill_total
                ) < 0.05

        return extracted_data

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@api_app.post("/api/export-excel")
async def export_nafeza_excel(data: InvoiceDataResponse):
    """Generates a Nafeza-compliant Excel document with structured headers, items, and lookup reference sheets."""
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Invoice"

        # Styling Definitions
        navy_header_fill = PatternFill(start_color="0A192F", end_color="0A192F", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        bold_font = Font(name="Calibri", size=10, bold=True)
        regular_font = Font(name="Calibri", size=10)
        
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        # -------------------------------------------------------------------
        # Header Metadata Block (Nafeza Custom Declaration Schema)
        # -------------------------------------------------------------------
        ws['B1'] = getattr(data, 'exporter_name', '') or "Foreign Supplier / Exporter"
        ws['B1'].font = bold_font
        
        ws['B3'] = "Commercial Invoice"
        ws['B3'].font = Font(name="Calibri", size=12, bold=True, color="1E3A8A")

        ws['B5'] = "Export To:"
        ws['C5'] = getattr(data, 'importer_name', '') or "Elamry Logistics Client"
        ws['B5'].font = bold_font

        ws['B7'] = "Egypt Tax Code :"
        ws['C7'] = getattr(data, 'importer_tax_id', '') or ""
        ws['D7'] = "ACID #:"
        ws['E7'] = getattr(data, 'acid_number', '') or ""
        ws['F7'] = "Origin Port:"
        ws['G7'] = getattr(data, 'origin_port', '') or ""

        ws['B8'] = "Purchase Order #:"
        ws['C8'] = getattr(data, 'purchase_order_num', '') or ""
        ws['D8'] = "Purchase Order Date:"
        ws['E8'] = getattr(data, 'purchase_order_date', '') or ""
        ws['F8'] = "Destination Port:"
        ws['G8'] = getattr(data, 'destination_port', '') or ""

        ws['B9'] = "Invoice #:"
        ws['C9'] = getattr(data, 'invoice_number', '') or ""
        ws['D9'] = "Invoice Date:"
        ws['E9'] = getattr(data, 'invoice_date', '') or ""

        ws['B10'] = "Currency:"
        ws['C10'] = getattr(data, 'currency', 'USD') or "USD"
        ws['D10'] = "Inco Term:"
        ws['E10'] = getattr(data, 'incoterm', 'FOB') or "FOB"

        for row in range(5, 11):
            for col in ['B', 'D', 'F']:
                cell = ws[f'{col}{row}']
                if cell.value:
                    cell.font = bold_font

        # -------------------------------------------------------------------
        # Line Items Table Header (Row 12)
        # -------------------------------------------------------------------
        table_headers = [
            "#", "Product Code", "TradeMarkOwner/Manufacturer", "Brand Name", 
            "Model", "HS Tariff Code", "Country of Origin", "Description", 
            "Quantity", "Qty Unit", "Expiry Date", "Unit Price", 
            "Gross Weight", "Net Weight", "Weight Unit", "Total"
        ]

        for col_idx, header in enumerate(table_headers, start=1):
            cell = ws.cell(row=12, column=col_idx, value=header)
            cell.fill = navy_header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # -------------------------------------------------------------------
        # Line Items Population (Starting Row 13)
        # -------------------------------------------------------------------
        current_row = 13
        items = data.items or []

        for idx, item in enumerate(items, start=1):
            ws.cell(row=current_row, column=1, value=idx).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=2, value=getattr(item, 'product_code', '') or f"ITEM-{idx}")
            ws.cell(row=current_row, column=3, value=getattr(data, 'exporter_name', '') or "")
            ws.cell(row=current_row, column=4, value=getattr(item, 'brand_name', '') or "")
            ws.cell(row=current_row, column=5, value=getattr(item, 'model', '') or "")
            ws.cell(row=current_row, column=6, value=getattr(item, 'hs_code', '') or "").alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=7, value=getattr(item, 'country_of_origin', 'CN') or "CN").alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=8, value=getattr(item, 'item_description', '') or "")
            
            qty_cell = ws.cell(row=current_row, column=9, value=getattr(item, 'quantity', 0) or 0)
            qty_cell.number_format = '#,##0.00'
            
            ws.cell(row=current_row, column=10, value=getattr(item, 'qty_unit', 'PCS') or "PCS").alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=11, value=getattr(item, 'expiry_date', '') or "")
            
            price_cell = ws.cell(row=current_row, column=12, value=getattr(item, 'unit_price', 0.0) or 0.0)
            price_cell.number_format = '#,##0.00'
            
            ws.cell(row=current_row, column=13, value=getattr(item, 'gross_weight', 0) or 0)
            ws.cell(row=current_row, column=14, value=getattr(item, 'net_weight', 0) or 0)
            ws.cell(row=current_row, column=15, value=getattr(item, 'weight_unit', 'KGM') or "KGM").alignment = Alignment(horizontal="center")
            
            # Excel formula for line item total calculation
            total_cell = ws.cell(row=current_row, column=16, value=f"=I{current_row}*L{current_row}")
            total_cell.number_format = '#,##0.00'
            total_cell.font = bold_font

            for col_idx in range(1, 17):
                ws.cell(row=current_row, column=col_idx).border = thin_border
                if col_idx not in [1, 6, 7, 10, 15, 16]:
                    ws.cell(row=current_row, column=col_idx).font = regular_font

            current_row += 1

        # Fallback if no items exist
        if not items:
            current_row = 13

        # -------------------------------------------------------------------
        # Totals & Reconciliation Summary Block
        # -------------------------------------------------------------------
        ws.cell(row=current_row, column=2, value="Lines #:").font = bold_font
        ws.cell(row=current_row, column=3, value=len(items)).font = bold_font
        ws.cell(row=current_row, column=15, value="Invoice Subtotal").font = bold_font
        ws.cell(row=current_row, column=16, value=f"=SUM(P13:P{current_row-1 if current_row > 13 else 13})").font = bold_font
        ws.cell(row=current_row, column=16).number_format = '#,##0.00'
        
        current_row += 1
        ws.cell(row=current_row, column=15, value="Freight Cost").font = bold_font
        ws.cell(row=current_row, column=16, value=getattr(data, 'freight_cost', 0.0) or 0.0)
        
        current_row += 1
        ws.cell(row=current_row, column=15, value="Insurance Cost").font = bold_font
        ws.cell(row=current_row, column=16, value=getattr(data, 'insurance_cost', 0.0) or 0.0)
        
        current_row += 1
        ws.cell(row=current_row, column=15, value="Other Costs").font = bold_font
        ws.cell(row=current_row, column=16, value=getattr(data, 'other_costs', 0.0) or 0.0)
        
        current_row += 1
        ws.cell(row=current_row, column=15, value="Total").font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
        ws.cell(row=current_row, column=16, value=f"=P{current_row-4}+P{current_row-3}+P{current_row-2}+P{current_row-1}").font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
        ws.cell(row=current_row, column=16).number_format = '#,##0.00'

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # -------------------------------------------------------------------
        # Nafeza Lookup Sheets (Units, Package Types, Currencies)
        # -------------------------------------------------------------------
        uom_ws = wb.create_sheet(title="UOM List")
        uom_ws.append(["UOMCode", "Unit Of Measure", "Currency", "Symbol"])
        uom_data = [
            ["KGM", "Kilogram", "EGP", "E£"],
            ["GRM", "Gram", "USD", "$"],
            ["STN", "Ton", "EUR", "€"],
            ["MTS", "Metric Ton", "GBP", "£"],
            ["PCS", "Piece", "CNY", "¥"],
            ["NAR", "Unit", None, None],
            ["MTR", "Metre", None, None],
        ]
        for row in uom_data:
            uom_ws.append(row)

        pkg_ws = wb.create_sheet(title="Package Type list")
        pkg_ws.append(["Code", "Name"])
        pkg_data = [
            ["BX", "Box"], ["CT", "Carton"], ["CR", "Crate"], 
            ["BD", "Board"], ["BG", "Bag"], ["RO", "Roll"], ["RL", "Reel"]
        ]
        for row in pkg_data:
            pkg_ws.append(row)

        # Save workbook to buffer
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"Nafeza_Invoice_{data.acid_number or 'Export'}.xlsx"
        headers = {'Content-Disposition': f'attachment; filename="{filename}"'}

        return Response(
            content=output.getvalue(),
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers=headers
        )

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate Excel file: {str(e)}")
web_app = Flask(__name__, template_folder="templates")

@web_app.route("/")
def index():
    """Renders the main dashboard interface from templates/index.html."""
    return render_template("index.html")

# Mount Flask inside FastAPI via WSGI Middleware
api_app.mount("/", WSGIMiddleware(web_app))

app = api_app

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)